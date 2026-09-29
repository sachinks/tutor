"""Build the tester workbook docs/testing/TUTOR-test-suite.xlsx from the sources of truth.

Sources (edit these, never the workbook's case text):
    docs/testing/test-cases/*.md        manual test cases
    docs/testing/automation-map.json    which automated tests cover each case
    docs/product/requirements.md        requirement IDs for traceability
    docs/testing/demo-data.md           demo accounts
    backend/apps/**/test*.py            automated tests (read with ast, no Django needed)
    backend/qa/smoke_test.py            smoke checks

Usage (from backend/):
    python qa/build_test_suite.py            # write the workbook
    python qa/build_test_suite.py --check    # only validate (CI): unknown test labels, malformed cases, duplicate IDs
"""

import argparse
import ast
import datetime as dt
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

BACKEND = Path(__file__).resolve().parent.parent
REPO = BACKEND.parent
DOCS = REPO / "docs"
CASES_DIR = DOCS / "testing" / "test-cases"
MAP_FILE = DOCS / "testing" / "automation-map.json"
REQS_FILE = DOCS / "product" / "requirements.md"
DEMO_FILE = DOCS / "testing" / "demo-data.md"
SMOKE_FILE = BACKEND / "qa" / "smoke_test.py"
OUTPUT = DOCS / "testing" / "TUTOR-test-suite.xlsx"

AREAS = {
    "accounts-and-consent": "Accounts & consent",
    "catalogue": "Catalogue",
    "learning-and-quizzes": "Learning & quizzes",
    "security-and-privacy": "Security & privacy",
    "admin": "Admin",
    "demo-and-hosted": "Demo world & hosted",
    "errors-and-logs": "Errors & logs",
    "ai-index": "AI lesson index",
    "ai-tutor": "AI tutor",
}
STATUSES = ["Not run", "Pass", "Fail", "Blocked", "N/A"]
HEADER_RE = re.compile(r"^\*\*(TC-[A-Z]+-\d+) · (.+?) · (P[1-3])(?: · (.+?))?\*\*\s*$")
REQ_RE = re.compile(r"\b(N?FR-[A-Z0-9]+-\d+)\b")


@dataclass
class Case:
    id: str
    area: str
    section: str
    title: str
    priority: str
    requirements: list
    pre: str = ""
    steps: str = ""
    expected: str = ""
    automated: list = field(default_factory=list)

    @property
    def environment(self):
        text = f"{self.pre} {self.steps}".lower()
        if self.id.startswith("TC-HOST"):
            return "Hosted"
        if "outbox" in text or "smoke_test.py`" in text and "--hosted" not in text:
            return "Local"
        if self.id.startswith("TC-DEMO"):
            return "Local or hosted"
        return "Local"


@dataclass
class AutoTest:
    label: str
    app: str
    layer: str
    cls: str
    name: str
    description: str
    covers: list = field(default_factory=list)


def _clean(text):
    return re.sub(r"\s+", " ", text.replace("**", "")).strip()


def _finish(case, body):
    """Split a case's body text into Pre / Steps / Expected."""
    parts = re.split(r"\b(Pre|Steps|Expected):", " ".join(body))
    fields = {"Pre": "", "Steps": "", "Expected": ""}
    for i in range(1, len(parts) - 1, 2):
        fields[parts[i]] = (fields[parts[i]] + " " + parts[i + 1]).strip()
    case.pre, case.steps, case.expected = (_clean(fields[k]) for k in ("Pre", "Steps", "Expected"))
    return case


def parse_cases():
    cases = []
    for path in sorted(CASES_DIR.glob("*.md")):
        area = AREAS.get(path.stem, path.stem)
        section, current, body = "", None, []
        for line in path.read_text(encoding="utf-8").splitlines():
            header = HEADER_RE.match(line.strip())
            if line.startswith("## ") or header:
                if current is not None:
                    cases.append(_finish(current, body))
                current, body = None, []
                if line.startswith("## "):
                    section = line[3:].strip()
                else:
                    reqs = REQ_RE.findall(header.group(4) or "")
                    current = Case(header.group(1), area, section, header.group(2), header.group(3), reqs)
            elif current is not None and line.strip():
                body.append(line.strip())
        if current is not None:
            cases.append(_finish(current, body))
    return cases


def _layer(path, cls_bases, name):
    if "test_queries" in path.name or "query" in name:
        return "Performance (query budget)"
    if "SimpleTestCase" in cls_bases:
        return "Unit"
    if path.parent.name == "demo":
        return "Demo data"
    if "test_api" in path.name or "test_security" in path.name:
        return "API"
    return "Integration"


def parse_automated():
    tests = []
    for path in sorted((BACKEND / "apps").rglob("test*.py")):
        if path.name == "testing.py":
            continue
        module = ".".join(path.relative_to(BACKEND).with_suffix("").parts)
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in tree.body:
            if not isinstance(node, ast.ClassDef):
                continue
            bases = " ".join(ast.unparse(b) for b in node.bases)
            for item in node.body:
                if isinstance(item, ast.FunctionDef) and item.name.startswith("test_"):
                    doc = ast.get_docstring(item) or item.name[5:].replace("_", " ").capitalize()
                    tests.append(
                        AutoTest(
                            label=f"{module}.{node.name}.{item.name}",
                            app=module.split(".")[1],
                            layer=_layer(path, bases, item.name),
                            cls=node.name,
                            name=item.name,
                            description=doc.splitlines()[0],
                        )
                    )
    return tests


def parse_smoke():
    source = SMOKE_FILE.read_text(encoding="utf-8")
    hosted_start = source.index("def hosted_checks")
    local_start = source.index("def main")
    checks = []
    for m in re.finditer(r'check\(\s*f?"([A-Z]\d+)\s+([^"]+)"', source):
        mode = "hosted" if hosted_start < m.start() < local_start else "smoke"
        checks.append((mode, m.group(1), m.group(2)))
    return checks


def parse_requirements():
    reqs = []
    for line in REQS_FILE.read_text(encoding="utf-8").splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) >= 3 and REQ_RE.fullmatch(cells[0]):
            if cells[0].startswith("NFR"):
                reqs.append((cells[0], cells[1], _clean(cells[2]), "—"))
            else:
                reqs.append((cells[0], cells[0].split("-")[1], _clean(cells[1]), _clean(cells[2])))
    return reqs


def parse_demo_accounts():
    rows, header = [], None
    for line in DEMO_FILE.read_text(encoding="utf-8").splitlines():
        if not line.startswith("|"):
            header = None
            continue
        cells = [_clean(c) for c in line.strip().strip("|").split("|")]
        if set("".join(cells)) <= set("-: "):
            continue
        if cells[0] == "Key":
            header = cells
            continue
        if header and "Mobile" in header:
            rec = dict(zip(header, cells, strict=False))
            rows.append(
                [
                    rec.get("Key", ""),
                    rec.get("Name", ""),
                    rec.get("Mobile", ""),
                    rec.get("Email", ""),
                    rec.get("Role / state") or rec.get("Role", ""),
                    rec.get("Use it to test") or f"Can: {rec.get('Can', '')}. Cannot: {rec.get('Cannot', '')}",
                ]
            )
    return rows


def validate(cases, tests, smoke, mapping):
    problems = []
    ids = [c.id for c in cases]
    for dup in sorted({i for i in ids if ids.count(i) > 1}):
        problems.append(f"duplicate case id {dup}")
    for c in cases:
        if not c.steps or not c.expected:
            problems.append(f"{c.id}: needs both 'Steps:' and 'Expected:'")
    known = {t.label for t in tests} | {f"{mode}:{cid}" for mode, cid, _ in smoke}
    for tc, labels in mapping.items():
        if tc.startswith("_"):
            continue
        if tc not in ids:
            problems.append(f"automation-map: unknown case {tc}")
        for label in labels:
            if label not in known:
                problems.append(f"automation-map: {tc} → unknown test {label}")
    return problems


def build_workbook(cases, tests, smoke, reqs, accounts, path):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    font = "Arial"
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(name=font, bold=True, color="FFFFFF")
    body_font = Font(name=font, size=10)
    input_fill = PatternFill("solid", fgColor="FFF7D6")  # cells testers fill in
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    def table(ws, headers, rows, widths, inputs=()):
        ws.append(headers)
        for col, _ in enumerate(headers, start=1):
            c = ws.cell(row=1, column=col)
            c.fill, c.font, c.alignment, c.border = (
                head_fill,
                head_font,
                Alignment(wrap_text=True, vertical="center"),
                border,
            )
        for row in rows:
            ws.append(row)
        for r in range(2, ws.max_row + 1):
            for col in range(1, len(headers) + 1):
                c = ws.cell(row=r, column=col)
                c.font, c.alignment, c.border = body_font, wrap, border
                if headers[col - 1] in inputs:
                    c.fill = input_fill
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 2)}"
        ws.row_dimensions[1].height = 30

    # --- Manual cases -----------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Manual cases"
    headers = [
        "ID",
        "Area",
        "Section",
        "Title",
        "Priority",
        "Requirements",
        "Environment",
        "Preconditions / accounts",
        "Steps",
        "Expected result",
        "Automated tests",
        "Automated by",
        "Assignee",
        "Status",
        "Tested on",
        "Build / commit",
        "Bug ID",
        "Notes",
    ]
    inputs = ("Assignee", "Status", "Tested on", "Build / commit", "Bug ID", "Notes")
    rows = [
        [
            c.id,
            c.area,
            c.section,
            c.title,
            c.priority,
            ", ".join(c.requirements),
            c.environment,
            c.pre,
            c.steps,
            c.expected,
            len(c.automated),
            "\n".join(c.automated),
            "",
            "Not run",
            "",
            "",
            "",
            "",
        ]
        for c in cases
    ]
    table(ws, headers, rows, [12, 18, 18, 34, 8, 16, 14, 34, 60, 60, 10, 50, 14, 11, 12, 14, 12, 30], inputs)
    last = ws.max_row
    status_dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    status_dv.error, status_dv.errorTitle = "Choose a status from the list.", "Status"
    ws.add_data_validation(status_dv)
    status_dv.add(f"N2:N{last}")
    date_dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
    ws.add_data_validation(date_dv)
    date_dv.add(f"O2:O{last}")
    for r in range(2, last + 1):
        ws.cell(row=r, column=15).number_format = "yyyy-mm-dd"
    colours = {"Pass": "D4EDDA", "Fail": "F8D7DA", "Blocked": "FFE8A1", "N/A": "E2E3E5"}
    for status, colour in colours.items():
        ws.conditional_formatting.add(
            f"A2:R{last}", FormulaRule(formula=[f'$N2="{status}"'], fill=PatternFill("solid", fgColor=colour))
        )
    manual_last = last

    # --- Automated tests --------------------------------------------------------------------------
    ws = wb.create_sheet("Automated tests")
    rows = [
        [
            t.label,
            t.app,
            t.layer,
            t.cls,
            t.description,
            ", ".join(t.covers),
            f"python manage.py test {t.label} --settings=config.settings.test",
        ]
        for t in tests
    ]
    table(
        ws,
        ["Test label", "App", "Layer", "Class", "What it proves", "Covers manual cases", "Run it (from backend/)"],
        rows,
        [70, 12, 22, 26, 60, 26, 80],
    )

    # --- Smoke checks -----------------------------------------------------------------------------
    ws = wb.create_sheet("Smoke checks")
    rows = [
        [
            "Local journey" if mode == "smoke" else "Hosted (read-only)",
            f"{mode}:{cid}",
            name,
            "python qa/smoke_test.py" if mode == "smoke" else "python qa/smoke_test.py --hosted --base <site URL>",
        ]
        for mode, cid, name in smoke
    ]
    table(ws, ["Mode", "Check", "What it checks", "Command (from backend/)"], rows, [20, 12, 60, 60])

    # --- Traceability -----------------------------------------------------------------------------
    ws = wb.create_sheet("Traceability")
    rows = []
    for rid, area, text, status in reqs:
        linked = [c for c in cases if rid in c.requirements]
        r = len(rows) + 2
        rows.append(
            [
                rid,
                area,
                text,
                status,
                ", ".join(c.id for c in linked),
                len(linked),
                sum(len(c.automated) for c in linked),
                f'=IF(F{r}=0,"No manual case",IF(G{r}=0,"Manual only","Manual + automated"))',
            ]
        )
    table(
        ws,
        [
            "Requirement",
            "Area",
            "Requirement text",
            "Build status",
            "Manual cases",
            "# manual",
            "# automated",
            "Coverage",
        ],
        rows,
        [14, 12, 60, 24, 40, 10, 12, 20],
    )
    trace_last = ws.max_row
    for text, colour in (("No manual case", "F8D7DA"), ("Manual only", "FFE8A1"), ("Manual + automated", "D4EDDA")):
        ws.conditional_formatting.add(
            f"H2:H{trace_last}", FormulaRule(formula=[f'$H2="{text}"'], fill=PatternFill("solid", fgColor=colour))
        )

    # --- Demo accounts ----------------------------------------------------------------------------
    ws = wb.create_sheet("Demo accounts")
    table(ws, ["Key", "Name", "Mobile", "Email", "Role / state", "Use it to test"], accounts, [16, 16, 14, 24, 40, 70])

    # --- Bug log ----------------------------------------------------------------------------------
    ws = wb.create_sheet("Bug log")
    example = [
        "BUG-001",
        "TC-QZ-06",
        "Answering twice returns 500 instead of 409",
        "High",
        "Local",
        "Start quiz; answer position 1; answer position 1 again",
        "409 conflict",
        "500 server error",
        "Tester name",
        dt.date(2026, 10, 1),
        "Open",
        "Example row: replace or delete",
    ]
    headers = [
        "Bug ID",
        "TC ID",
        "Title",
        "Severity",
        "Environment",
        "Steps to reproduce",
        "Expected",
        "Actual",
        "Reported by",
        "Reported on",
        "Status",
        "Notes",
    ]
    table(ws, headers, [example], [10, 12, 40, 10, 14, 50, 30, 30, 16, 12, 12, 30], inputs=tuple(headers))
    for dv_range, options in (
        ("D2:D500", "Critical,High,Medium,Low"),
        ("K2:K500", "Open,In progress,Fixed,Verified,Won't fix"),
        ("E2:E500", "Local,Hosted"),
    ):
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(dv_range)
    ws["J2"].number_format = "yyyy-mm-dd"

    # --- Read me (first sheet) ----------------------------------------------------------------------
    ws = wb.create_sheet("Read me", 0)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 70
    title = ws.cell(row=1, column=1, value="TUTOR test suite")
    title.font = Font(name=font, size=16, bold=True, color="1F3A5F")
    ws.cell(
        row=2, column=1, value=f"Generated {dt.date.today():%Y-%m-%d} by backend/qa/build_test_suite.py"
    ).font = Font(name=font, size=9, italic=True, color="666666")
    m = f"'Manual cases'!$N$2:$N${manual_last}"
    summary = [
        ("Manual cases", f"=COUNTA('Manual cases'!$A$2:$A${manual_last})", "Written in docs/testing/test-cases/*.md"),
        ("  Priority P1", f"=COUNTIF('Manual cases'!$E$2:$E${manual_last},\"P1\")", "Must pass before any release"),
        ("  Priority P2", f"=COUNTIF('Manual cases'!$E$2:$E${manual_last},\"P2\")", ""),
        ("  Priority P3", f"=COUNTIF('Manual cases'!$E$2:$E${manual_last},\"P3\")", ""),
        (
            "  With automated coverage",
            f"=COUNTIF('Manual cases'!$K$2:$K${manual_last},\">0\")",
            "Automated tests exist too",
        ),
        ("  Automated coverage %", "=IF(B4=0,0,B8/B4)", "Share of manual cases also covered by automation"),
        ("Status: Pass", f'=COUNTIF({m},"Pass")', "Filled in by testers in 'Manual cases'"),
        ("Status: Fail", f'=COUNTIF({m},"Fail")', "Every Fail needs a Bug ID"),
        ("Status: Blocked", f'=COUNTIF({m},"Blocked")', "Can't be run (environment, missing feature)"),
        ("Status: Not run", f'=COUNTIF({m},"Not run")', ""),
        ("Pass rate (of cases run)", "=IF((B10+B11+B12)=0,0,B10/(B10+B11+B12))", "Pass ÷ (Pass + Fail + Blocked)"),
        (
            "Automated tests",
            f"=COUNTA('Automated tests'!$A$2:$A${len(tests) + 1})",
            "Django tests; run by ./dev.sh and CI",
        ),
        ("Smoke checks", f"=COUNTA('Smoke checks'!$A$2:$A${len(smoke) + 1})", "Black-box checks over HTTP"),
        (
            "Requirements without a manual case",
            f'=COUNTIF(Traceability!$H$2:$H${trace_last},"No manual case")',
            "Planned features show here until they are built",
        ),
    ]
    ws.cell(row=3, column=1, value="Summary").font = Font(name=font, bold=True, size=12)
    for i, (label, formula, note) in enumerate(summary, start=4):
        ws.cell(row=i, column=1, value=label).font = body_font
        cell = ws.cell(row=i, column=2, value=formula)
        cell.font = Font(name=font, size=10, bold=True)
        if "%" in label or "rate" in label:
            cell.number_format = "0.0%"
        ws.cell(row=i, column=3, value=note).font = Font(name=font, size=9, color="666666")
    row = 4 + len(summary) + 1
    guide = [
        ("How to use", ""),
        ("1", "Filter 'Manual cases' by Area or Priority, put your name in Assignee (yellow cells are yours to fill)."),
        ("2", "Run the case exactly as written. Accounts are in 'Demo accounts' and docs/testing/demo-data.md."),
        (
            "3",
            "Set Status, Tested on and Build / commit. For every Fail, add a row in 'Bug log' "
            "and put its ID in Bug ID.",
        ),
        ("4", "Don't edit case text here: change docs/testing/test-cases/*.md in a pull request, then rebuild."),
        ("Statuses", "Not run · Pass · Fail · Blocked · N/A"),
        (
            "Environments",
            "Local = your machine with tester tools (dev outbox). Hosted = https://tutor-platform-ovlg.onrender.com",
        ),
        (
            "Severity",
            "Critical: data leak, security or consent bypass · High: a P1 case fails · "
            "Medium: a P2 case fails · Low: cosmetic",
        ),
        ("Rebuild", "From backend/: python qa/build_test_suite.py (also validates the automation map)"),
    ]
    for label, text in guide:
        a = ws.cell(row=row, column=1, value=label)
        a.font = Font(
            name=font, bold=label == "How to use" or not label.isdigit(), size=12 if label == "How to use" else 10
        )
        c = ws.cell(row=row, column=3, value=text)
        c.font, c.alignment = body_font, Alignment(wrap_text=True, vertical="top")
        row += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--check", action="store_true", help="validate only; don't write the workbook")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args(argv)

    cases, tests, smoke = parse_cases(), parse_automated(), parse_smoke()
    mapping = json.loads(MAP_FILE.read_text(encoding="utf-8"))
    problems = validate(cases, tests, smoke, mapping)
    by_label = {t.label: t for t in tests}
    for c in cases:
        c.automated = list(mapping.get(c.id, []))
        for label in c.automated:
            if label in by_label:
                by_label[label].covers.append(c.id)
    if problems:
        print("Test suite sources have problems:", *problems, sep="\n  ")
        return 1
    print(f"{len(cases)} manual cases, {len(tests)} automated tests, {len(smoke)} smoke checks: OK")
    if not args.check:
        build_workbook(cases, tests, smoke, parse_requirements(), parse_demo_accounts(), args.output)
        print(f"Wrote {args.output.relative_to(REPO) if args.output.is_relative_to(REPO) else args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
